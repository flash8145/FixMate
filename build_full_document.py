import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_full_document(filename):
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        # Header / Footer setup
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("FixMate — System Specification & Requirement Analysis Document")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("FixMate: Hyperlocal Blue-Collar Service Marketplace | Confidential & Academic Project Report")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Styling helper functions
    def set_cell_background(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def add_callout_box(title, text, bg_hex="F8FAFC", border_hex="EA580C"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.5)
        
        cell = tbl.cell(0, 0)
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run_t = p.add_run(f"{title}\n")
        run_t.bold = True
        run_t.font.name = 'Calibri'
        run_t.font.size = Pt(11)
        run_t.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)
        
        run_b = p.add_run(text)
        run_b.font.name = 'Calibri'
        run_b.font.size = Pt(10)
        run_b.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(4)

    def style_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(22)
        h.paragraph_format.space_after = Pt(8)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = 'Calibri'
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(0x0F, 0x29, 0x42) # Deep Navy
        return h

    def style_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = 'Calibri'
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Royal Blue
        return h

    def style_h3(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = 'Calibri'
        r.font.size = Pt(11.5)
        r.font.color.rgb = RGBColor(0xEA, 0x58, 0x0C) # Dark Warm Orange
        return h

    def add_p(text, bold_prefix=None, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
            r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        return p

    def add_bullet(text, bold_title=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_title:
            r_bold = p.add_run(bold_title)
            r_bold.bold = True
            r_bold.font.name = 'Calibri'
            r_bold.font.size = Pt(10.5)
            r_bold.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        return p

    def render_table(col_widths, headers, data):
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        # Header row
        hdr_row = table.rows[0]
        for idx, heading in enumerate(headers):
            cell = hdr_row.cells[idx]
            cell.text = heading
            set_cell_background(cell, "0F2942") # Deep Navy
            set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.bold = True
                r.font.name = 'Calibri'
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
        # Data rows
        for r_idx, row_data in enumerate(data):
            row = table.add_row()
            bg_hex = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.text = str(cell_value)
                set_cell_background(cell, bg_hex)
                set_cell_margins(cell, top=90, bottom=90, left=130, right=130)
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                    
        # Apply widths
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width
        
        # Spacer paragraph after table
        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(0)
        p_spacer.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # DOCUMENT TITLE / COVER BLOCK
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    r_main = title_p.add_run("FixMate — Hyperlocal Blue-Collar Service Marketplace")
    r_main.bold = True
    r_main.font.name = 'Calibri'
    r_main.font.size = Pt(22)
    r_main.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(16)
    r_sub = subtitle_p.add_run("System Architecture, Project Specification & Comprehensive Requirement Analysis Report")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(0xEA, 0x58, 0x0C)
    
    add_callout_box(
        "Executive Summary",
        "FixMate is a modern full-stack MERN (MongoDB, Express.js, React 18, Node.js) web application engineered to bridge the gap between local residents and nearby blue-collar service professionals (plumbers, electricians, carpenters, AC technicians, painters). Utilizing an interactive map-first architecture with Leaflet and OpenStreetMap, FixMate delivers geospatial worker discovery, upfront transparent dynamic cost estimation, a deterministic finite state machine booking lifecycle, and privacy-shielded direct communication without middleman commissions.",
        bg_hex="F1F5F9",
        border_hex="0F2942"
    )

    # -------------------------------------------------------------
    # SECTION 1: INTRODUCTION
    # -------------------------------------------------------------
    style_h1("1. INTRODUCTION")
    
    # 1.1 Existing System
    style_h2("1.1 Existing System")
    add_p(
        "In the contemporary urban landscape, securing reliable, skilled, and prompt blue-collar services for household maintenance—such as plumbing, electrical troubleshooting, carpentry, air conditioning servicing, and painting—remains largely fragmented and inefficient. The existing ecosystem primarily relies on three traditional operational models:"
    )
    add_bullet(
        " Relies entirely on informal recommendations from neighbors, security guards, or local hardware store shopkeepers. This system lacks standardization, offers zero transparency regarding the technician's actual skill level, experience, or track record, and provides no guarantee of availability or punctuality.",
        "1. Informal Word-of-Mouth & Physical Search:"
    )
    add_bullet(
        " Legacy telephone directory platforms provide static contact lists of technicians. However, these lists fail to indicate real-time technician availability, do not support geospatial proximity visualization, and frequently result in customer harassment through unsolicited lead-sharing where customer contact numbers are auctioned off to dozens of independent contractors simultaneously.",
        "2. Classified Listings and Lead-Generation Directories (e.g., JustDial, Sulekha):"
    )
    add_bullet(
        " Corporate aggregator platforms act as rigid middlemen between customers and service providers. While they offer standard booking interfaces, they operate on exploitative commission structures (deducting 20% to 35% from the blue-collar worker's earnings), strictly conceal worker identity until auto-dispatch occurs, mandate fixed catalog pricing that ignores job complexity, and restrict direct human negotiation.",
        "3. Centralized Corporate Aggregator Agencies (e.g., Urban Company):"
    )
    add_p(
        "In summary, the existing solutions force consumers to choose between an unverified, risky informal market and an expensive, restrictive corporate middleman that disempowers grassroots service professionals."
    )

    # 1.2 Need for the New System
    style_h2("1.2 Need for the New System")
    add_p(
        "There is an urgent requirement for a decentralized, transparent, and map-first service marketplace that combines the visual convenience of modern ride-hailing applications (such as Uber or Rapido) with the autonomy and direct peer-to-peer communication of a classified marketplace. The fundamental imperatives driving the development of FixMate include:"
    )
    add_bullet(
        " Consumers must be able to visually see active, available service technicians situated around their exact geographic coordinates or chosen neighborhood on an interactive map, eliminating the guesswork associated with distant service dispatch.",
        "• Hyperlocal Visual Proximity:"
    )
    add_bullet(
        " Unlike monopolistic aggregator platforms that impose heavy commission cuts on blue-collar laborers, FixMate facilitates direct matching where workers retain 100% of their earnings, settling payments directly with customers offline via cash or UPI.",
        "• Elimination of Intermediary Commission:"
    )
    add_bullet(
        " Providing customers with immediate, server-calculated cost estimation ranges prior to booking confirmation, accounting for worker visit charges, category base rates, travel distance surcharges, urgency premiums, and nighttime multipliers.",
        "• Upfront Price Transparency:"
    )
    add_bullet(
        " Empowering blue-collar tradespeople to manage their own digital presence, set custom hourly/visit rates, define their exact service coverage radius in kilometers, and toggle their real-time availability on or off at will.",
        "• Blue-Collar Worker Autonomy:"
    )
    add_bullet(
        " Preventing lead-spamming and unwanted solicitations by strictly shielding worker phone numbers and addresses until a booking is explicitly accepted by the technician.",
        "• Data Privacy & Contact Gating:"
    )

    # 1.3 Objectives of the New System
    style_h2("1.3 Objectives of the New System")
    add_p(
        "The primary objective of FixMate is to design, implement, and deploy a robust, high-performance, and accessible web-based platform connecting residents with local service providers. Specific measurable objectives include:"
    )
    add_bullet(
        " Implement MongoDB 2dsphere indexing and an aggregation pipeline to execute sub-100ms geospatial proximity queries, matching workers based on distance, service radius, active status, and verification.",
        "1. Geospatial Discovery Engine:"
    )
    add_bullet(
        " Deliver a frictionless frontend using React 18, Vite, and Leaflet (OpenStreetMap) allowing users to search via GPS geolocation or debounced Nominatim city lookup, view interactive map pins, and inspect comprehensive worker profiles.",
        "2. Interactive Map-First Interface:"
    )
    add_bullet(
        " Develop an isolated, pure-function pricing engine that computes transparent, deterministic price ranges from admin-calibrated parameters without hardcoded business constants.",
        "3. Algorithmic Pricing Engine:"
    )
    add_bullet(
        " Enforce a rigorous state machine (PENDING → ACCEPTED → IN_PROGRESS → PRICE_SET → COMPLETED / REJECTED / CANCELLED) with an immutable audit timeline capturing every status change, timestamp, and actor role.",
        "4. Deterministic Booking State Machine:"
    )
    add_bullet(
        " Facilitate frictionless offline collaboration by unlocking technician phone numbers and pre-filled WhatsApp deep links immediately upon booking acceptance.",
        "5. Direct Contact & Communication Gateway:"
    )
    add_bullet(
        " Establish a verified feedback loop where only customers with COMPLETED bookings can submit ratings and reviews, with atomic denormalized rating aggregation on worker profiles.",
        "6. Verified Reputation System:"
    )
    add_bullet(
        " Provide a centralized administration panel for category management (CRUD), worker verification toggle, and live pricing engine calibration.",
        "7. Comprehensive Platform Administration:"
    )

    # 1.4 Problem Definition
    style_h2("1.4 Problem Definition")
    add_p(
        "The urban household maintenance sector suffers from acute information asymmetry, high transaction friction, and economic exploitation. The problem can be formally decomposed into three key perspectives:"
    )
    
    add_p(
        "When unexpected plumbing leaks, short circuits, or appliance breakdowns occur, residents struggle to quickly locate trustworthy nearby technicians. Customers face uncertain wait times, lack visibility into fair local market pricing, experience anxiety over post-service overcharging, and risk privacy infringement from aggressive telemarketing directories.",
        bold_prefix="A. The Customer Problem: "
    )
    add_p(
        "Local blue-collar tradespeople possess valuable vocational skills but lack affordable digital tools to market their services locally. When working through corporate aggregator agencies, they are subjected to punishing commission fees, penalized for declining distant jobs, stripped of direct customer rapport, and denied control over their work schedules and pricing.",
        bold_prefix="B. The Service Worker Problem: "
    )
    add_p(
        "Traditional digital solutions either treat blue-collar services as static yellow-page phone lists (ignoring real-time availability and location) or as fully centralized taxi-dispatch operations (ignoring the necessity for custom problem assessment, variable scoping, and direct client negotiation).",
        bold_prefix="C. The Architectural & Market Problem: "
    )

    # 1.5 Core Definition
    style_h2("1.5 Core Definition")
    add_p(
        "The foundational architectural concepts, domain entities, and core mechanics governing the FixMate system are defined as follows:"
    )
    add_bullet(
        " The core spatial discovery mechanism utilizing MongoDB's 2dsphere spherical geometry index. Workers are indexed by GeoJSON Point coordinates [longitude, latitude]. The pipeline executes a two-stage filter: first querying candidates within a global distance cap via $geoNear, followed by an in-memory evaluation matching the customer's distance against each worker's individual service radius (serviceRadiusM).",
        "• Hyperlocal Geospatial Indexing & Matching:"
    )
    add_bullet(
        " A dynamic, customizable boundary (stored in meters, configured in kilometers) set by each worker defining the maximum radial distance they are willing to travel from their registered base location.",
        "• Service Coverage Radius Model:"
    )
    add_bullet(
        " A deterministic mathematical calculation engine that computes estimated service cost ranges as a pure function: Total Estimate = (Visit Charge + Category Base Rate + Travel Charge + Urgency Fee) × Night Multiplier ± Spread Percentage. Operating without database side-effects, it derives parameters dynamically from an admin-editable pricing configuration.",
        "• Pure Pricing Engine:"
    )
    add_bullet(
        " A formal finite state machine governing the life cycle of every service request. Transitions are constrained by strict role permissions (Customer, Worker, Admin) and append to an immutable audit timeline array.",
        "• Booking State Machine & Audit Timeline:"
    )
    add_bullet(
        " A security model ensuring that sensitive technician contact data (phone numbers and WhatsApp triggers) remains redacted at the API projection level until the booking status reaches ACCEPTED.",
        "• Contact Gating & Role-Based Access Control (RBAC):"
    )
    add_bullet(
        " A database optimization strategy where worker rating statistics (sum, count, average) are updated atomically via $inc operators upon review creation, preventing expensive O(n) runtime aggregations during spatial search queries.",
        "• Atomic Denormalized Rating Aggregation:"
    )

    # 1.6 Project Profile
    style_h2("1.6 Project Profile")
    add_p(
        "The following table provides a comprehensive technical and operational summary of the FixMate project profile:"
    )
    
    headers_profile = ["Attribute", "Project Specification / Profile Details"]
    data_profile = [
        ["Project Title", "FixMate — Hyperlocal Blue-Collar Service Marketplace"],
        ["Project Type", "Web Application / Academic Final Year Capstone Project"],
        ["Target Industry", "On-Demand Home Services & Blue-Collar Gig Economy"],
        ["Architecture Pattern", "3-Tier Layered Architecture (Controller → Service → Model)"],
        ["Frontend Framework", "React 18 (Vite SPA) with Zustand State Management"],
        ["UI Styling System", "Tailwind CSS with Responsive Mobile-First Design"],
        ["Mapping & Geocoding", "Leaflet (`react-leaflet`), OpenStreetMap, Nominatim API"],
        ["Backend Runtime", "Node.js (v18+ LTS) with Express.js RESTful API"],
        ["Database Layer", "MongoDB Atlas with Mongoose ODM & Native `2dsphere` Geospatial Indexing"],
        ["Authentication & RBAC", "Stateless JSON Web Tokens (JWT), Bcrypt Password Hashing, Custom RBAC Middleware"],
        ["Pricing Model", "Server-Side Dynamic Multi-Factor Pure Pricing Engine"],
        ["Payment Method", "Direct Offline Settlement (Cash / Direct Peer-to-Peer UPI)"],
        ["Communication", "Gated Direct Phone Contact + Prefilled WhatsApp Deep Links"]
    ]
    render_table([Inches(2.2), Inches(4.3)], headers_profile, data_profile)

    # 1.7 Assumptions and Constraints
    style_h2("1.7 Assumptions and Constraints")
    
    style_h3("1.7.1 System Assumptions")
    add_bullet(
        " All customers and technicians have access to modern internet-connected devices (smartphones, tablets, or desktop computers) equipped with standard modern web browsers.",
        "1. Device & Network Connectivity:"
    )
    add_bullet(
        " Customers grant browser geolocation access for automatic coordinate detection or utilize the built-in Nominatim search to specify their location accurately.",
        "2. Geolocation Accessibility:"
    )
    add_bullet(
        " Workers maintain active mobile phone numbers capable of receiving direct cellular phone calls and WhatsApp messages.",
        "3. Communication Channels:"
    )
    add_bullet(
        " Customers and technicians mutually agree to settle the final service fee offline upon job completion via physical currency or peer-to-peer digital UPI payments.",
        "4. Financial Settlement:"
    )
    add_bullet(
        " Technicians honestly configure their primary base coordinates, skill rates, and realistic travel radiuses.",
        "5. Worker Authenticity:"
    )

    style_h3("1.7.2 System Constraints")
    add_bullet(
        " To eliminate third-party billing dependencies (e.g., Google Maps API billing cards), the platform relies on OpenStreetMap and Leaflet. Nominatim geocoding enforces a strict 1 request/second rate-limit, necessitating debounced frontend queries.",
        "1. Zero-Cost Open Source Map API Constraint:"
    )
    add_bullet(
        " In alignment with the academic project scope, native mobile apps, automated payment gateway escrow, live background worker GPS telemetry tracking, and automated biometric KYC verification are deliberately excluded.",
        "2. Bounded Academic Scope:"
    )
    add_bullet(
        " Coordinate data must adhere strictly to the GeoJSON convention: [longitude, latitude]. Inverting coordinates will cause spatial query failures.",
        "3. Geospatial Coordinate Ordering Constraint:"
    )
    add_bullet(
        " Token-based authentication requires clients to store JWT tokens securely and attach them as Bearer tokens in HTTP Authorization headers for protected endpoints.",
        "4. Stateless Authentication Constraint:"
    )

    # 1.8 Advantages and Limitations of the Proposed System
    style_h2("1.8 Advantages and Limitations of the Proposed System")
    
    style_h3("1.8.1 Advantages of FixMate")
    add_bullet(
        " Unlike traditional aggregators that obscure worker identity, FixMate gives customers full autonomy to browse, filter, inspect ratings, and select their preferred technician directly from the map.",
        "• Transparent Visual Selection:"
    )
    add_bullet(
        " Zero platform commission ensures blue-collar workers retain their complete earnings while customers avoid inflated platform surcharge fees.",
        "• 100% Worker Revenue Retention:"
    )
    add_bullet(
        " Algorithmic multi-factor ranking balances distance, rating averages, and completed job volume to present the best technicians first.",
        "• Weighted Proximity Scoring:"
    )
    add_bullet(
        " Upfront calculation of minimum and maximum price ranges eliminates billing surprises before work commences.",
        "• Predictable Cost Estimation:"
    )
    add_bullet(
        " Built entirely on open-source libraries without proprietary recurring API subscription costs.",
        "• Cost-Effective & Scalable Architecture:"
    )

    style_h3("1.8.2 Limitations of the Proposed System")
    add_bullet(
        " Because payments are transacted directly between customer and technician, the system relies on user confirmation rather than automated bank webhook settlement.",
        "• Offline Payment Verification:"
    )
    add_bullet(
        " Worker location represents their registered home/workshop base rather than live turn-by-turn vehicle tracking.",
        "• Static Base Location vs. Live Telemetry:"
    )
    add_bullet(
        " In-app messaging is replaced by direct WhatsApp integration, requiring external app switching for text communication.",
        "• External Messaging Dependency:"
    )

    # -------------------------------------------------------------
    # SECTION 2: REQUIREMENT DETERMINATION & ANALYSIS
    # -------------------------------------------------------------
    doc.add_page_break()
    style_h1("2. REQUIREMENT DETERMINATION & ANALYSIS")
    
    # 2.1 Requirement Determination
    style_h2("2.1 Requirement Determination")
    add_p(
        "Requirement determination is the systematic process of identifying, analyzing, and documenting the functional capabilities, performance standards, architectural constraints, and system specifications necessary to deliver a fully operational FixMate platform."
    )

    style_h3("2.1.1 Functional Requirements Specification")
    add_p(
        "Functional requirements define the core software operations, calculations, business workflows, and user interactions implemented across each subsystem:"
    )

    headers_fr = ["Req ID", "Module / Area", "Functional Requirement Description"]
    data_fr = [
        ["FR-C01", "Customer", "User registration and login with phone, email, password, and customer role assignment."],
        ["FR-C02", "Customer", "Interactive geospatial worker search with category filtering, radius slider, and city lookup."],
        ["FR-C03", "Customer", "Worker detail exploration showing bio, experience, skills, visit fees, rating average, and reviews."],
        ["FR-C04", "Customer", "Dynamic price estimate retrieval prior to booking submission based on location and schedule."],
        ["FR-C05", "Customer", "Booking creation with problem description, image upload, schedule time, and urgency level."],
        ["FR-C06", "Customer", "Real-time access to worker phone number and WhatsApp deep link once status is ACCEPTED."],
        ["FR-C07", "Customer", "Final price review and confirmation upon technician job completion."],
        ["FR-C08", "Customer", "Verified review submission (1-5 star rating, feedback text, tags) for COMPLETED bookings."],
        ["FR-W01", "Worker", "Worker registration, authentication, and comprehensive profile setup (bio, experience, skills)."],
        ["FR-W02", "Worker", "Service rate definition (base rate per visit / per hour) linked to specific service categories."],
        ["FR-W03", "Worker", "Geographic base point pinpointing on interactive map and service radius configuration (km)."],
        ["FR-W04", "Worker", "Real-time availability status toggle (Online / Offline) controlling search visibility."],
        ["FR-W05", "Worker", "Incoming booking notification queue with capability to Accept or Reject pending requests."],
        ["FR-W06", "Worker", "Customer contact details and address revelation upon booking acceptance."],
        ["FR-W07", "Worker", "Booking status advancement to IN_PROGRESS and final price submission upon job completion."],
        ["FR-W08", "Worker", "Worker performance dashboard tracking completed jobs, rating aggregates, and active bookings."],
        ["FR-A01", "Admin", "Service category management (Create, Read, Update, Deactivate service categories and icons)."],
        ["FR-A02", "Admin", "Worker verification management (inspecting worker profiles and toggling verified badges)."],
        ["FR-A03", "Admin", "Dynamic pricing configuration management (per-km rate, free distance, urgency fee, night multipliers)."],
        ["FR-A04", "Admin", "Platform analytics and statistics aggregation (user counts, active bookings, completed jobs)."],
        ["FR-S01", "System Core", "Two-stage MongoDB geospatial pipeline ($geoNear + $expr radius check) with weighted ranking."],
        ["FR-S02", "System Core", "Pure-function pricing engine calculating transparent estimate ranges without DB side-effects."],
        ["FR-S03", "System Core", "Strict booking state machine enforcing allowed role transitions and immutable audit timeline logging."],
        ["FR-S04", "System Core", "Atomic rating aggregation ($inc updates on worker profile) upon review creation."],
        ["FR-S05", "System Core", "Strict API-level contact gating redacting phone numbers prior to ACCEPTED state."],
        ["FR-S06", "System Core", "Centralized error handling middleware with typed AppError responses and input validation."]
    ]
    render_table([Inches(1.1), Inches(1.5), Inches(3.9)], headers_fr, data_fr)

    style_h3("2.1.2 Non-Functional Requirements Specification")
    add_p(
        "Non-functional requirements specify the quality attributes, architectural qualities, security standards, and performance benchmarks governing the FixMate system:"
    )

    headers_nfr = ["Quality Attribute", "Specification & Benchmark Standard"]
    data_nfr = [
        ["Performance & Latency", "Geospatial search queries must execute in < 100ms on indexed MongoDB collections. Frontend map pan/zoom and city search queries must be debounced (300ms–500ms) to prevent excessive server requests."],
        ["Security & Data Protection", "Passwords must be hashed using bcrypt with at least 10 salt rounds. Authentication must use stateless JWT tokens with at least 32-character secrets. Worker phone numbers must be stripped from public search DTOs."],
        ["Scalability & Efficiency", "Denormalized rating aggregates must ensure O(1) rating reads during search. Database indexes on location, category, availability, and city ensure optimal pipeline execution at scale."],
        ["Reliability & Consistency", "State machine transitions must be atomic (using findOneAndUpdate) to prevent race conditions during concurrent worker acceptance. Centralized error handling prevents unhandled crashes."],
        ["Usability & Responsiveness", "Fully responsive UI adapting seamlessly across mobile, tablet, and desktop viewports using Tailwind CSS. Intuitive map pins with visual status indicators and clear error alerts."],
        ["Maintainability & Modularity", "Strict layered architecture (Controller → Service → Model). Controllers never touch Mongoose directly; business logic is encapsulated in unit-testable service modules."]
    ]
    render_table([Inches(2.2), Inches(4.3)], headers_nfr, data_nfr)

    style_h3("2.1.3 Hardware and Software Requirements")
    add_p(
        "The development, deployment, and operational execution of FixMate necessitate the following hardware and software infrastructure:"
    )

    headers_env = ["Component", "Minimum Requirement", "Recommended Specification"]
    data_dev_hw = [
        ["Processor (CPU)", "Dual-Core 2.0 GHz (x86/x64)", "Intel Core i5 / AMD Ryzen 5 (4+ Cores)"],
        ["System Memory (RAM)", "4 GB DDR4", "8 GB – 16 GB DDR4/DDR5"],
        ["Storage Space", "2 GB Free Storage (HDD/SSD)", "10 GB Free SSD Storage"],
        ["Network Adapter", "Standard 10 Mbps Internet Connection", "High-Speed Broadband / 4G / 5G"]
    ]
    add_p("Table 2.1: Development & Server Hardware Requirements", bold_prefix="")
    render_table([Inches(1.8), Inches(2.2), Inches(2.5)], headers_env, data_dev_hw)

    data_client_hw = [
        ["Mobile / Tablet Devices", "Android 8.0+ or iOS 12+ with GPS", "Modern Smartphone with 4G/5G & Chrome/Safari"],
        ["Desktop / Laptop", "Any standard PC with web browser", "Display Resolution 1366x768 or higher"],
        ["Input Peripherals", "Touchscreen, Mouse, Keyboard", "Touchscreen or Pointer Device"]
    ]
    add_p("Table 2.2: Client Hardware Requirements", bold_prefix="")
    render_table([Inches(1.8), Inches(2.2), Inches(2.5)], headers_env, data_client_hw)

    headers_sw = ["Layer / Domain", "Technology & Version Specification", "Role in FixMate"]
    data_sw = [
        ["Operating System", "Windows 10/11, Ubuntu Linux 20.04+, macOS", "Host development & server operating environment"],
        ["JavaScript Runtime", "Node.js v18.x / v20.x LTS", "Backend asynchronous server execution runtime"],
        ["Package Manager", "npm v9.x or higher", "Dependency management for server and client"],
        ["Database Engine", "MongoDB Atlas / MongoDB Server v6.0+", "NoSQL database with 2dsphere geospatial indexing"],
        ["Backend Framework", "Express.js v4.19+", "RESTful API routing and middleware framework"],
        ["Frontend Framework", "React 18 (Vite Bundler)", "Reactive single-page user interface application"],
        ["State Management", "Zustand v4.5+", "Lightweight client-side authentication state store"],
        ["Map Library", "Leaflet v1.9+ (`react-leaflet`)", "Interactive map rendering and custom pin markers"],
        ["Geocoding Service", "Nominatim (OpenStreetMap API)", "City and locality search coordinate conversion"],
        ["CSS Framework", "Tailwind CSS v3.4+", "Utility-first responsive styling and layout system"],
        ["Authentication", "JSON Web Token (`jsonwebtoken`), `bcryptjs`", "Stateless token security and password hashing"],
        ["Development Tools", "VS Code, Postman, Git, Chrome DevTools", "Code authoring, API testing, version control"]
    ]
    add_p("Table 2.3: Software Stack & Environment Specifications", bold_prefix="")
    render_table([Inches(1.8), Inches(2.3), Inches(2.4)], headers_sw, data_sw)

    # 2.2 Targeted User
    style_h2("2.2 Targeted User")
    add_p(
        "Target user analysis identifies the primary personas, demographic profiles, operational contexts, and behavioral expectations of the distinct stakeholder groups interacting with the FixMate platform."
    )

    style_h3("2.2.1 Primary User Roles and Personas")
    
    add_callout_box(
        "Persona 1: Urban Resident / Homeowner (Customer)",
        "• Name / Profile: Priya Sharma, 32, Urban Apartment Resident & Working Professional.\n"
        "• Scenario: Experiencing a sudden pipe leakage under the kitchen sink on a weekday evening.\n"
        "• Needs & Goals: Wants to quickly open the app, locate a licensed plumber within 3 km on the map, check verified customer reviews, review an upfront price estimate range, book immediately, and contact the technician via WhatsApp without paying high commission markups.\n"
        "• Frustrations: Aggressive calls from telemarketing agencies, lack of price transparency, and unresponsive service providers.",
        bg_hex="F8FAFC",
        border_hex="0F2942"
    )

    add_callout_box(
        "Persona 2: Independent Skilled Tradesperson (Service Worker / Technician)",
        "• Name / Profile: Ramesh Kumar, 41, Independent Electrician & Appliance Technician.\n"
        "• Scenario: Skilled blue-collar technician operating independently in a designated urban cluster.\n"
        "• Needs & Goals: Wants a simple mobile web interface to showcase his experience, define his service radius (5 km), set fair base inspection rates, toggle his availability online when free, receive direct customer bookings, and collect 100% of his payment directly in cash or UPI.\n"
        "• Frustrations: High 25-30% commission deductions by corporate aggregator apps, forced distant travel, and lack of direct customer relationship.",
        bg_hex="FFF7ED",
        border_hex="EA580C"
    )

    add_callout_box(
        "Persona 3: System Administrator / Platform Moderator (Admin)",
        "• Name / Profile: Platform Operations Specialist.\n"
        "• Scenario: Managing platform integrity, service standards, and operational parameters.\n"
        "• Needs & Goals: Needs an administrative console to add/edit service categories, review and verify new technician profiles, adjust dynamic pricing parameters (per-km travel fee, night multiplier, spread percentage), and monitor platform metrics.\n"
        "• Frustrations: Hardcoded system constants and lack of granular operational control.",
        bg_hex="F1F5F9",
        border_hex="2563EB"
    )

    style_h3("2.2.2 User Characteristics, Demographics & Accessibility Considerations")
    add_p(
        "To ensure maximum adoption and usability across diverse user segments, FixMate is engineered around the following audience characteristics:"
    )
    add_bullet(
        " While urban customers possess high digital fluency, blue-collar workers often operate with entry-level smartphones and varying literacy levels. FixMate addresses this through high-contrast visual cues, intuitive map pins, prominent status toggles, and direct integration with familiar communication tools like WhatsApp and standard cellular calling.",
        "1. Varied Digital Literacy Spectrum:"
    )
    add_bullet(
        " The platform is tailored for Tier-1, Tier-2, and Tier-3 urban and semi-urban clusters where residential neighborhoods, apartment complexes, and commercial zones require localized maintenance support within compact 3 to 15-kilometer radiuses.",
        "2. Hyperlocal Geographic Focus:"
    )
    add_bullet(
        " Customers prioritize speed and cost predictability during emergencies, while workers prioritize geographical proximity and immediate direct payment. FixMate balances both priorities through transparent upfront estimates and direct offline financial settlement.",
        "3. Behavioral Dynamics & Trust:"
    )
    add_bullet(
        " The interface strictly adheres to responsive design principles, ensuring seamless touch navigation, readable typography, and optimal performance on both low-cost mobile devices and high-resolution desktop screens.",
        "4. Device & Accessibility Optimization:"
    )

    # Save document
    doc.save(filename)
    print(f"Document successfully created and saved to: {filename}")

if __name__ == "__main__":
    out_path = r"d:\Fixmate\FixMate_Project_Documentation_Section1_and_2.docx"
    create_full_document(out_path)
