import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def style_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.bold = True
    r.font.name = 'Calibri'
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)
    return h

def style_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.bold = True
    r.font.name = 'Calibri'
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    return h

def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_data_table(doc, headers, data, col_widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    # Header Row
    hdr_row = table.rows[0]
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.text = heading
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.bold = True
            r.font.name = 'Calibri'
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Data Rows
    for r_idx, row_data in enumerate(data):
        row = table.add_row()
        bg_hex = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = str(cell_value)
            set_cell_background(cell, bg_hex)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.1
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                if c_idx == 0:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Set column widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Add space after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(10)

def generate_document():
    doc = docx.Document()
    
    # Set standard page margins (0.75 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    headers = ["Field Name", "Data Type", "Constraints / Keys", "Description", "Sample Data"]
    col_widths = [Inches(1.4), Inches(0.9), Inches(1.3), Inches(1.8), Inches(1.6)]

    style_heading_1(doc, "4. Data Dictionary")
    add_body(doc, "The FixMate database consists of 6 primary collections implemented in MongoDB Atlas using Mongoose ODM. Below are the structural definitions, field data types, constraints, descriptions, and sample values for each entity.")

    # 1. users
    style_heading_2(doc, "Table 4.1: users Collection")
    add_body(doc, "Stores authentication credentials, contact details, profile avatars, and permanent home addresses for all user roles (Customers, Workers, and Administrators).")
    users_data = [
        ["_id", "ObjectId", "Primary Key", "Unique identifier for the user account.", "66a1b2c3d4e5f6789012345a"],
        ["name", "String", "Required, Trimmed", "Full name of the user / registered username.", "Ramesh Sharma"],
        ["phone", "String", "Required, Unique, 10 Digits", "Primary mobile number used for login and WhatsApp coordination.", "9820098200"],
        ["email", "String", "Optional, Lowercase", "Email address of the user.", "ramesh@fixmate.in"],
        ["passwordHash", "String", "Required, Hidden", "Bcrypt-hashed password string (salt cost 10).", "$2a$10$wT8Kj...e9qQ"],
        ["role", "String", "Required, Enum: ['customer', 'worker', 'admin']", "Platform authorization role and permission tier.", "worker"],
        ["avatarUrl", "String", "Optional, Default: \"\"", "Image URL or preset identifier for user display photo.", "https://images.unsplash.com/..."],
        ["address.line", "String", "Optional, Default: \"\"", "House number, flat, street, or building name.", "Flat 304, Green Heights"],
        ["address.area", "String", "Optional, Default: \"\"", "Locality, neighborhood, or sector.", "Kothrud"],
        ["address.city", "String", "Optional, Default: \"\"", "City name.", "Pune"],
        ["address.pincode", "String", "Optional, Default: \"\"", "6-digit postal code.", "411038"],
        ["address.location.type", "String", "Enum: ['Point'], Default: 'Point'", "GeoJSON type identifier for geospatial indexing.", "Point"],
        ["address.location.coordinates", "Array [Number]", "[Longitude, Latitude]", "Geographic coordinates of permanent address.", "[73.8567, 18.5204]"],
        ["isActive", "Boolean", "Default: true", "Account active state flag.", "true"],
        ["createdAt", "Date", "Auto-generated", "Account registration timestamp.", "2026-08-15T09:30:00.000Z"],
    ]
    add_data_table(doc, headers, users_data, col_widths)

    # 2. workerprofiles
    style_heading_2(doc, "Table 4.2: workerprofiles Collection")
    add_body(doc, "Contains professional credentials, service offerings, workshop location, travel charges, and rating statistics for service workers.")
    workers_data = [
        ["_id", "ObjectId", "Primary Key", "Unique identifier for the worker profile.", "66a1b2c3d4e5f6789012345b"],
        ["userId", "ObjectId", "Required, Unique, FK (users._id)", "Reference to the associated user account.", "66a1b2c3d4e5f6789012345a"],
        ["bio", "String", "Optional, Default: \"\"", "Professional summary and work description.", "Experienced plumber with 8+ years expertise."],
        ["experienceYrs", "Number", "Default: 0, Min: 0", "Years of trade experience.", "8"],
        ["skills", "Array [Object]", "Required, Non-empty", "List of skills and associated rates offered.", "[ { categoryId: '...', baseRate: 350 } ]"],
        ["skills.categoryId", "ObjectId", "Required, FK (categories._id)", "Reference to service category.", "66a1b2c3d4e5f6789012345c"],
        ["skills.baseRate", "Number", "Required, Min: 0", "Base labor charge (in ₹) for this specific skill.", "350"],
        ["skills.rateUnit", "String", "Enum: ['PER_VISIT', 'PER_HOUR']", "Charging model (per visit vs. hourly labor).", "PER_VISIT"],
        ["visitCharge", "Number", "Default: 0, Min: 0", "Minimum inspection / call-out fee (in ₹).", "150"],
        ["perKmRate", "Number", "Default: 0, Min: 0", "Custom travel charge (in ₹) per km beyond free distance.", "20"],
        ["location.type", "String", "Enum: ['Point'], Default: 'Point'", "GeoJSON Point indexed with 2dsphere.", "Point"],
        ["location.coordinates", "Array [Number]", "Required, [Lng, Lat]", "Base workshop coordinates for map distance calculations.", "[73.8540, 18.5180]"],
        ["address.line", "String", "Optional", "Base workshop street address.", "Shop 12, Guruprasad Complex"],
        ["address.area", "String", "Optional", "Base workshop locality.", "Kothrud"],
        ["address.city", "String", "Optional", "Base workshop city.", "Pune"],
        ["address.pincode", "String", "Optional", "Base workshop postal code.", "411038"],
        ["serviceRadiusM", "Number", "Default: 5000 (5 km)", "Service coverage radius in meters.", "10000"],
        ["isAvailable", "Boolean", "Default: false", "Online/Offline toggle for live map visibility.", "true"],
        ["isVerified", "Boolean", "Default: false", "Verification badge approval status granted by Admin.", "true"],
        ["ratingAgg.sum", "Number", "Default: 0", "Sum of all star ratings received.", "96"],
        ["ratingAgg.count", "Number", "Default: 0", "Total number of reviews received.", "20"],
        ["ratingAgg.avg", "Number", "Default: 0.0, Range: 0.0 - 5.0", "Atomic average rating (sum / count).", "4.8"],
        ["jobsCompleted", "Number", "Default: 0", "Total completed bookings count.", "34"],
        ["photos", "Array [String]", "Default: []", "URLs of past work portfolio or certification photos.", "[\"https://fixmate.in/w1.jpg\"]"],
        ["createdAt", "Date", "Auto-generated", "Profile creation timestamp.", "2026-08-16T11:00:00.000Z"],
    ]
    add_data_table(doc, headers, workers_data, col_widths)

    # 3. categories
    style_heading_2(doc, "Table 4.3: categories Collection")
    add_body(doc, "Stores standardized service categories available on the platform.")
    categories_data = [
        ["_id", "ObjectId", "Primary Key", "Unique identifier for the category.", "66a1b2c3d4e5f6789012345c"],
        ["slug", "String", "Required, Unique, Lowercase", "URL slug identifier.", "plumber"],
        ["name", "String", "Required, Trimmed", "Display name of the service category.", "Plumber"],
        ["icon", "String", "Optional, Default: \"\"", "Lucide icon identifier or emoji for UI display.", "Wrench"],
        ["isActive", "Boolean", "Default: true", "Active status toggle managed by admin.", "true"],
    ]
    add_data_table(doc, headers, categories_data, col_widths)

    # 4. bookings
    style_heading_2(doc, "Table 4.4: bookings Collection")
    add_body(doc, "Records service requests, upfront estimates, status lifecycle transitions, and final bills.")
    bookings_data = [
        ["_id", "ObjectId", "Primary Key", "Unique identifier for the booking.", "66a1b2c3d4e5f6789012345d"],
        ["code", "String", "Required, Unique, Uppercase", "Human-readable reference code.", "BK-84291"],
        ["customerId", "ObjectId", "Required, FK (users._id)", "Reference to customer who booked.", "66a1b2c3d4e5f6789012345e"],
        ["workerId", "ObjectId", "Required, FK (workerprofiles._id)", "Reference to selected worker profile.", "66a1b2c3d4e5f6789012345b"],
        ["categoryId", "ObjectId", "Required, FK (categories._id)", "Reference to booked service category.", "66a1b2c3d4e5f6789012345c"],
        ["problemText", "String", "Required, Trimmed", "Description of repair issue entered by customer.", "Kitchen sink pipe leakage"],
        ["photos", "Array [String]", "Default: []", "Optional photo uploads showing damage.", "[\"https://fixmate.in/leak.jpg\"]"],
        ["address.line", "String", "Required, Min: 3 chars", "Destination street address for job execution.", "A-502, Rohan Tarang"],
        ["address.city", "String", "Required, Min: 2 chars", "Destination city name.", "Pune"],
        ["address.location.type", "String", "Enum: ['Point'], Default: 'Point'", "GeoJSON Point for destination location.", "Point"],
        ["address.location.coordinates", "Array [Number]", "Required, [Lng, Lat]", "Job site coordinates for map directions.", "[73.8495, 18.5132]"],
        ["scheduledAt", "Date", "Required", "Scheduled service appointment date and time.", "2026-08-20T14:30:00.000Z"],
        ["urgency", "String", "Enum: ['NORMAL', 'IMMEDIATE']", "Urgency level defining surcharge applicability.", "NORMAL"],
        ["estimate.breakdown", "Object", "Nested Document", "Breakdown of visit fee, travel charge, and subtotal.", "{ visitCharge: 150, subtotal: 540 }"],
        ["estimate.min", "Number", "Required, Min: 0", "Lower bound of upfront estimate range (-15%).", "460"],
        ["estimate.max", "Number", "Required, Min: 0", "Upper bound of upfront estimate range (+15%).", "620"],
        ["final.amount", "Number", "Optional", "Final agreed bill submitted by worker upon completion.", "550"],
        ["final.setBy", "ObjectId", "FK (users._id)", "User reference who submitted final bill.", "66a1b2c3d4e5f6789012345a"],
        ["final.confirmedAt", "Date", "Optional", "Timestamp when customer confirmed final bill.", "2026-08-20T16:15:00.000Z"],
        ["status", "String", "Enum: ['PENDING', 'ACCEPTED', 'IN_PROGRESS', 'PRICE_SET', 'COMPLETED', ...]", "Current lifecycle stage of the booking.", "COMPLETED"],
        ["timeline", "Array [Object]", "Audit Log", "History of status transitions (status, at, by).", "[{ status: 'COMPLETED' }]"],
        ["createdAt", "Date", "Auto-generated", "Booking submission timestamp.", "2026-08-20T13:45:00.000Z"],
        ["updatedAt", "Date", "Auto-generated", "Last modification timestamp.", "2026-08-20T16:15:00.000Z"],
    ]
    add_data_table(doc, headers, bookings_data, col_widths)

    # 5. reviews
    style_heading_2(doc, "Table 4.5: reviews Collection")
    add_body(doc, "Stores customer ratings, feedback tags, and reviews after a job is completed.")
    reviews_data = [
        ["_id", "ObjectId", "Primary Key", "Unique identifier for the review.", "66a1b2c3d4e5f6789012345f"],
        ["bookingId", "ObjectId", "Required, Unique, FK (bookings._id)", "Enforces strictly one review per completed booking.", "66a1b2c3d4e5f6789012345d"],
        ["workerId", "ObjectId", "Required, FK (workerprofiles._id)", "Reference to worker receiving the rating.", "66a1b2c3d4e5f6789012345b"],
        ["customerId", "ObjectId", "Required, FK (users._id)", "Reference to customer writing the review.", "66a1b2c3d4e5f6789012345e"],
        ["rating", "Number", "Required, Integer, Min: 1, Max: 5", "Star rating score.", "5"],
        ["tags", "Array [String]", "Default: []", "Selected feedback tags.", "[\"Punctual\", \"Clean Work\", \"Polite\"]"],
        ["text", "String", "Optional, Trimmed", "Written review commentary.", "Fixed the pipe leakage quickly!"],
        ["createdAt", "Date", "Auto-generated", "Review submission timestamp.", "2026-08-20T16:30:00.000Z"],
    ]
    add_data_table(doc, headers, reviews_data, col_widths)

    # 6. pricingconfigs
    style_heading_2(doc, "Table 4.6: pricingconfigs Collection")
    add_body(doc, "Global parameters controlling the server-side dynamic pricing calculation engine.")
    pricing_data = [
        ["_id", "ObjectId", "Primary Key", "Unique identifier for configuration document.", "66a1b2c3d4e5f67890123450"],
        ["perKmRate", "Number", "Required, Default: 15", "Default travel rate (in ₹) per km beyond free distance.", "15"],
        ["freeDistanceM", "Number", "Required, Default: 1000", "Free distance threshold in meters before mileage charges apply.", "1000"],
        ["urgencyFee", "Number", "Required, Default: 100", "Fixed emergency surcharge (in ₹) for IMMEDIATE jobs.", "100"],
        ["nightMultiplier", "Number", "Required, Default: 1.25", "Multiplier (+25%) applied to night scheduled bookings.", "1.25"],
        ["nightStartHour", "Number", "Required, Default: 21 (9:00 PM)", "Start hour (24h clock) when night tariff applies.", "21"],
        ["nightEndHour", "Number", "Required, Default: 6 (6:00 AM)", "End hour (24h clock) when night tariff expires.", "6"],
        ["estimateSpreadPct", "Number", "Required, Default: 0.15 (+-15%)", "Percentage spread around estimate creating [min, max] range.", "0.15"],
        ["deviationThresholdPct", "Number", "Required, Default: 0.20 (20%)", "Threshold above estimate before final bill triggers price alert.", "0.20"],
        ["updatedAt", "Date", "Auto-generated", "Timestamp when admin last updated pricing parameters.", "2026-08-01T10:00:00.000Z"],
    ]
    add_data_table(doc, headers, pricing_data, col_widths)

    doc.save("D:\\Fixmate\\FixMate_Data_Dictionary.docx")
    print("SUCCESS: FixMate_Data_Dictionary.docx generated successfully!")

if __name__ == "__main__":
    generate_document()
